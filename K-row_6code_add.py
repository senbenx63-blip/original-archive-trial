import openpyxl
import re

# 1. Excelファイルのロード
file_path = "dictionary_final.xlsx"  # ここにご自身のファイル名を指定してください
wb = openpyxl.load_workbook(file_path)

# 全シートを対象に処理
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    
    # --- 【A1セルの処理】年2桁の取得 ---
    a1_val = ws["A1"].value
    if a1_val is None:
        continue  # A1セルが空の場合はスキップ
        
    a1_str = str(a1_val)
    # 数字のみを抽出し、その下2桁を取得 (例: "2025年" や 2025 -> "25")
    year_digits = re.findall(r'\d+', a1_str)
    if not year_digits:
        continue  # A1から年が取得できない場合はスキップ
        
    year_full = year_digits[0]
    year_2digits = year_full[-2:].zfill(2) # 下2桁（足りない場合は0埋め）

    # --- シートごとの重複管理用辞書 ---
    code_counts = {}

    # --- 【行ごとの処理】 ---
    # ★ 1行目から最終行までループするように修正しました
    for row in range(1, ws.max_row + 1):
        f_val = ws.cell(row=row, column=6).value   # F列 (6列目)
        k_val = ws.cell(row=row, column=11).value  # K列 (11列目)

        # F列が空欄ではない、かつ K列が空欄 の行のみ処理
        if f_val is not None and str(f_val).strip() != "":
            if k_val is None or str(k_val).strip() == "":
                
                f_str = str(f_val)
                
                # かっこ内の「月/日」パターンを正規表現で探す (半角・全角かっこ両対応)
                match = re.search(r'[（\(].*?(\d{1,2})/(\d{1,2}).*?[）\)]', f_str)

                if match:
                    month = match.group(1).zfill(2) # 月を2桁化 (例: 1 -> 01)
                    day = match.group(2).zfill(2)   # 日を2桁化 (例: 8 -> 08)
                    
                    # ベースとなる「月日4桁 ＋ 年2桁」の6桁文字列を作成
                    base_code = f"{month}{day}{year_2digits}"
                    
                    # 出現回数をカウント＆連番付与のロジック
                    if base_code in code_counts:
                        code_counts[base_code] += 1
                        # 2回目以降は -2, -3 を末尾に追加
                        final_code = f"{base_code}-{code_counts[base_code]}"
                    else:
                        code_counts[base_code] = 1
                        # 1回目はそのまま6桁コードを使用
                        final_code = base_code
                    
                    # K列に【文字列形式】で書き込み
                    cell_k = ws.cell(row=row, column=11)
                    cell_k.number_format = '@'  # セルの表示形式を文字列に設定
                    cell_k.value = str(final_code)

# 2. 上書き保存
wb.save(file_path)
wb.close()

print("すべてのシートに対する処理と上書き保存が完了しました。")