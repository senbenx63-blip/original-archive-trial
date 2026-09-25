import re
import sqlite3
from pathlib import Path
import openpyxl

DB_PATH = "chama_archive.db"
XLSX_PATH = Path(__file__).parent / "dictionary_final.xlsx"

D_PARAM_RE = re.compile(r"[?&]d=([0-9\-]+)")
DATE_IN_PAREN_RE = re.compile(r"\((?:[^\d/()]*)?(\d{1,2})/(\d{1,2})(?:[^\d/()]*)?\)")
YOUTUBE_ID_RE = re.compile(r"[?&]v=([^&]+)")


def migrate_schema(conn: sqlite3.Connection):
    """streamsテーブルにyoutube関連の列が無ければ追加する"""
    cols = {row[1] for row in conn.execute("PRAGMA table_info(streams)").fetchall()}
    if "youtube_video_id" not in cols:
        conn.execute("ALTER TABLE streams ADD COLUMN youtube_video_id TEXT")
    if "youtube_url" not in cols:
        conn.execute("ALTER TABLE streams ADD COLUMN youtube_url TEXT")


def ensure_stream_stub(conn: sqlite3.Connection, stream_key: str):
    """まだstreamsに無いstream_keyなら、器だけ作っておく(stream_dateは自動計算)"""
    conn.execute(
        "INSERT INTO streams (stream_key) VALUES (?) ON CONFLICT(stream_key) DO NOTHING",
        (stream_key,),
    )


def compute_key_from_date(day_label, year):
    """F列の '1日目(1/9)' のような文字列とシートのA1セルの年号から MMDDYY を作る"""
    if day_label is None or year is None:
        return None
    m = DATE_IN_PAREN_RE.search(str(day_label))
    if not m:
        return None
    month, day = int(m.group(1)), int(m.group(2))
    yy = int(year) % 100
    return f"{month:02d}{day:02d}{yy:02d}"


def resolve_stream_key(row, year, seen_counts: dict):
    """K列 → G列URLのd=パラメータ → F列+シート年号 の優先順位でstream_keyを決める

    seen_countsは今回の実行中に「F列+年から計算したベースキー」が
    何回出てきたかを数える辞書。DBの中身を見るのではなく、
    Excelの読み込み順だけで判定するので、何度再実行しても同じ結果になる。
    """
    key_cell = row[10] if len(row) > 10 else None   # K列
    url_cell = row[6] if len(row) > 6 else None      # G列
    day_label = row[5] if len(row) > 5 else None     # F列

    if key_cell:
        return str(key_cell).strip()

    if url_cell:
        m = D_PARAM_RE.search(str(url_cell))
        if m:
            return m.group(1)

    # F列+年からの計算だけは、同じ日に2つ配信があると同じキーになりうるので
    # 「今回の実行で何回目の出現か」を数えて -2, -3... を自動で付ける
    base_key = compute_key_from_date(day_label, year)
    if base_key is None:
        return None

    seen_counts[base_key] = seen_counts.get(base_key, 0) + 1
    n = seen_counts[base_key]
    return base_key if n == 1 else f"{base_key}-{n}"


def main():
    conn = sqlite3.connect(DB_PATH)
    migrate_schema(conn)

    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    count = 0
    skipped = 0
    seen_counts = {}  # F列+年から計算したベースキーの出現回数(ブック全体で共有)

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        year = ws.cell(row=1, column=1).value  # シートの一番左上のセル(A1)= 年号
        current_title = None  # タイトルが空の行に、直前のタイトルを引き継ぐための変数

        for row in ws.iter_rows(values_only=True):
            title_cell = row[1] if len(row) > 1 else None  # B列
            url_cell = row[6] if len(row) > 6 else None     # G列

            if title_cell:
                current_title = title_cell

            stream_key = resolve_stream_key(row, year, seen_counts)
            if not stream_key:
                skipped += 1
                continue

            youtube_id = None
            youtube_url = None
            if url_cell:
                m = YOUTUBE_ID_RE.search(str(url_cell))
                if m:
                    youtube_id = m.group(1)
                    youtube_url = f"https://www.youtube.com/watch?v={youtube_id}"

            ensure_stream_stub(conn, stream_key)
            conn.execute(
                """UPDATE streams SET
                     title = COALESCE(title, ?),
                     youtube_video_id = COALESCE(?, youtube_video_id),
                     youtube_url = COALESCE(?, youtube_url)
                   WHERE stream_key = ?""",
                (current_title, youtube_id, youtube_url, stream_key),
            )
            count += 1

    conn.commit()
    conn.close()
    print(f"{count} 件の配信情報を辞書から反映しました(未特定でスキップ: {skipped}件)")


if __name__ == "__main__":
    main()